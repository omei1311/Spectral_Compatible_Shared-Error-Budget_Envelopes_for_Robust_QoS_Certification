"""Build the English V4.1 manuscript with native editable Word equations."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUTPUT=ROOT/'Spectral_Compatible_Shared_Error_Budget_QoS_Certification_V4.1_EN.docx'


def el(name, **attrs):
    e=OxmlElement('m:'+name)
    for k,v in attrs.items():e.set(qn('m:'+k),str(v))
    return e


def mr(text, plain=False):
    r=el('r')
    if plain:
        rp=el('rPr');rp.append(el('nor'));r.append(rp)
    t=el('t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t)
    return r


def container(name,nodes):
    out=el(name)
    for n in nodes:out.append(n)
    return out


def split_top(s,sep):
    depth=0;parts=[];start=0
    for i,c in enumerate(s):
        if c in '{([':depth+=1
        if c in '})]':depth-=1
        if c==sep and depth==0:parts.append(s[start:i]);start=i+1
    parts.append(s[start:]);return parts


class MathParser:
    def __init__(self,text):
        self.s=text.replace('\\{','{').replace('\\}','}')
        self.i=0
        self.last_close=None
    def brace(self):
        assert self.s[self.i]=='{',self.s[self.i:]
        self.i+=1;start=self.i;depth=1
        while depth:
            c=self.s[self.i];self.i+=1
            depth+=(c=='{')-(c=='}')
        return self.s[start:self.i-1]
    def argument(self, word=False):
        if self.s[self.i]=='{':return MathParser(self.brace()).parse()
        if word:
            token=re.match(r'[A-Za-z]+|[0-9]+',self.s[self.i:])
            if token:
                self.i+=len(token[0]);return MathParser(token[0]).parse()
        if self.s.startswith('∥',self.i):
            self.i+=1;return [mr('∥')]
        c=self.s[self.i];self.i+=1;return [mr(c)]
    def parse(self,closing=None):
        nodes=[]
        while self.i<len(self.s):
            c=self.s[self.i]
            if closing is not None and c in ')]}':
                self.i+=1;self.last_close=c;break
            keyword=next((k for k in ['frac','mat','cases','diag','text'] if self.s.startswith(k+'{',self.i)),None)
            if keyword:
                self.i+=len(keyword);content=self.brace()
                if keyword=='frac':
                    den=self.brace();f=el('f');f.append(container('num',MathParser(content).parse()));f.append(container('den',MathParser(den).parse()));nodes.append(f)
                elif keyword in ('mat','cases'):
                    matrix=el('m')
                    for row in split_top(content,';'):
                        r=el('mr')
                        for cell in split_top(row,'&'):r.append(container('e',MathParser(cell.strip()).parse()))
                        matrix.append(r)
                    d=el('d');pr=el('dPr');pr.append(el('begChr',val='(' if keyword=='mat' else '{'));pr.append(el('endChr',val=')' if keyword=='mat' else ''));d.append(pr);d.append(container('e',[matrix]));nodes.append(d)
                elif keyword=='diag':
                    nodes.append(mr('diag',True));nodes.append(self.delimited(MathParser(content).parse(),'(',')'))
                else:nodes.append(mr(content,True))
                continue
            if c=='√':
                self.i+=1;rad=el('rad');pr=el('radPr');pr.append(el('degHide',val=1));rad.append(pr);rad.append(el('deg'));rad.append(container('e',self.argument()));nodes.append(rad);continue
            if c in '_^':
                self.i+=1
                if not nodes:raise ValueError('script missing base: '+self.s)
                base=nodes.pop();args=self.argument(word=c=='_')
                scripted=el('sSub' if c=='_' else 'sSup');scripted.append(container('e',[base]));scripted.append(container('sub' if c=='_' else 'sup',args));nodes.append(scripted);continue
            if c in '([{':
                self.i+=1;end={'(':')','[':']','{':'}'}[c]
                content=self.parse(end)
                nodes.append(self.delimited(content,c,self.last_close or end));continue
            if c in ')]}':raise ValueError('unmatched bracket: '+self.s)
            operator=next((w for w in ('max','min','sin','cos','Re','inf') if self.s.startswith(w,self.i)),None)
            if operator:
                nodes.append(mr(operator,True));self.i+=len(operator);continue
            nodes.append(mr(c));self.i+=1
        return nodes
    @staticmethod
    def delimited(nodes,beg,end):
        d=el('d');pr=el('dPr');pr.append(el('begChr',val=beg));pr.append(el('endChr',val=end));d.append(pr);d.append(container('e',nodes));return d


def math(text):return container('oMath',MathParser(text).parse())


def add_inline(p,text):
    pieces=re.split(r'(\$[^$]+\$)',text)
    for part in pieces:
        if part.startswith('$') and part.endswith('$'):p._p.append(math(part[1:-1]))
        elif part:
            for j,chunk in enumerate(re.split(r'(\*\*[^*]+\*\*)',part)):
                if not chunk:continue
                run=p.add_run(chunk[2:-2] if chunk.startswith('**') else chunk)
                if chunk.startswith('**'):run.bold=True


def table(doc,caption,headers,rows,widths):
    p=doc.add_paragraph(caption,style='Caption');p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,h in enumerate(headers):t.rows[0].cells[i].text=h
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):cells[i].text=str(val)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:b.set(qn('w:'+k),v)
        borders.append(b)
    t._tbl.tblPr.append(borders)
    for rowidx,row in enumerate(t.rows):
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side,val in [('top','65'),('bottom','65'),('left','65'),('right','65')]:
                elem=OxmlElement('w:'+side);elem.set(qn('w:w'),val);elem.set(qn('w:type'),'dxa');marg.append(elem)
            tcpr.append(marg)
            if rowidx==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');tcpr.append(shade)
            for p in cell.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
                for run in p.runs:run.font.size=Pt(9);run.bold=rowidx==0
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def add_table(doc,name):
    if name=='boundary':
        table(doc,'Table 1. Explicit boundary checks with a = 1 throughout.',
              ['p','a','I₀','J⊥','c∥','Q∂','Reference q*'],
              [['2','1','1','0','1','0.590000','0.590000'],['0.2','1','1','0','−0.8','−1.202000','−1.202000'],
               ['2','1','1.3','0.830662','1','−0.181020','−0.181020'],['0.2','1','1.3','0.830662','−0.8','−1.963482','−1.963482']],
              [.45,.4,.5,.8,.5,1.2,1.42])
    elif name=='sensitivity':
        rows=[['0.5','0.06',118,2,117,118,118,118,0],['0.5','0.15',108,12,91,102,103,102,1],['0.5','0.30',63,57,6,26,28,21,2],
              ['1','0.06',90,30,78,81,83,90,2],['1','0.15',56,64,21,34,36,49,2],['1','0.30',7,113,0,0,0,0,0],
              ['2','0.06',24,96,13,18,18,22,0],['2','0.15',6,114,0,1,4,6,3],['2','0.30',0,120,0,0,0,0,0],
              ['Total','',472,608,326,380,390,408,10]]
        table(doc,'Table 2. Directly accepted units on the complete sensitivity grid, 256 cells.',
              ['γ','ρ','Feas.','Infeas.','D','Old Q','Q_G','L_quad','Added'],rows,[.45,.5,.62,.68,.45,.68,.65,.82,.52])
    elif name=='resolution':
        table(doc,'Table 3. Interval resolution on the same 1080 candidate units.',
              ['Cells','Old Q accepts','Q_G accepts','Additional accepts','Feasible units'],
              [[256,380,390,10,472],[1024,380,390,10,472],[4096,381,390,9,472]],[.6,1.25,1.25,1.5,.67])
    elif name=='timing':
        table(doc,'Table 4. End to end seconds; total reference calls across ten seeds in parentheses.',
              ['γ','Reference','D + fallback','Old Q + fallback','Q_G + fallback','L_quad + fallback'],
              [['0.5','0.1787\n(1321)','0.0375\n(244)','0.0313\n(160)','0.0316\n(147)','0.0248\n(164)'],
               ['1','0.1396\n(1005)','0.0564\n(379)','0.0516\n(306)','0.0511\n(297)','0.0429\n(285)'],
               ['2','0.0874\n(620)','0.0607\n(410)','0.0612\n(384)','0.0640\n(377)','0.0536\n(353)']],
              [.35,1.,1.,1.,1.,.92])


def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.top_margin=sec.bottom_margin=Inches(.85);sec.left_margin=sec.right_margin=Inches(1)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',19),('Heading 1',13),('Heading 2',11.5),('Caption',9.5)]:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        if name!='Caption':style.font.bold=True
        style.paragraph_format.space_before=Pt(10 if name.startswith('Heading') else 0)
        style.paragraph_format.space_after=Pt(5)
    equation=doc.styles.add_style('Equation',1);equation.base_style=normal
    equation.font.name='Cambria Math';equation.font.size=Pt(10.5)
    equation.paragraph_format.space_before=Pt(3);equation.paragraph_format.space_after=Pt(6)
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(6.2),WD_ALIGN_PARAGRAPH.RIGHT)
    # Document-wide math font and size inherited from the paragraph.
    settings=doc.settings.element;mp=el('mathPr');mp.append(el('mathFont',val='Cambria Math'));settings.append(mp)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    for raw in (ROOT/'paper_v41.md').read_text(encoding='utf-8').splitlines():
        line=raw.strip()
        if not line:continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title');continue
        if line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2');continue
        if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1');continue
        if line.startswith('EQ '):
            label,expr=line[3:].split(':',1);p=doc.add_paragraph(style='Equation')
            p._p.append(math(expr.strip().rstrip('.')));p.add_run('\t'+label)
            continue
        if line.startswith('TABLE '):add_table(doc,line[6:]);continue
        if line=='FIGURE spectral':
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ROOT/'results_review/figures/fig2_spectral_scan.png'),width=Inches(6.2))
            doc.add_paragraph('Figure 1. Constructed spectral scan. Approximate envelopes are diagnostic; the dashed interval lower bound is used for certification. The family changes the nominal channel norm.','Caption')
            continue
        p=doc.add_paragraph();add_inline(p,line)
    doc.core_properties.title='Spectral Compatible Shared Error Budget Envelopes for Robust QoS Certification'
    doc.core_properties.subject='English V4.1 with repaired nondegenerate and boundary sharpness proofs'
    doc.core_properties.version='4.1'
    for tree in (doc.styles.element,doc.element):
        for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
        for fonts in tree.iter(qn('w:rFonts')):
            for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                fonts.attrib.pop(qn('w:'+attr),None)
    for name in ('Title','Heading 1','Heading 2','Caption'):
        doc.styles[name].font.name='Times New Roman'
    ppr=doc.paragraphs[0]._p.get_or_add_pPr();no_border=OxmlElement('w:pBdr')
    for side in ('top','left','bottom','right'):
        child=OxmlElement('w:'+side);child.set(qn('w:val'),'nil');no_border.append(child)
    ppr.append(no_border)
    doc.save(OUTPUT)
    print(OUTPUT)
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(OUTPUT) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        print('native equations',len(xml.findall('.//m:oMath',ns)),'tables',len(xml.findall('.//w:tbl',ns)))


if __name__=='__main__':main()
