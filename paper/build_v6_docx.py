"""Build the English V6 manuscript with native editable Word equations.

Same rendering pipeline as V5 (which reuses V4.1); table inventory updated
for V6 (11 tables incl. related work, realistic pool, controlled scaling).
"""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_v41_docx import add_inline, el, math, table

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
OUTPUT = ROOT / 'Sharp_Robust_QoS_Certification_from_Six-Scalar_Spectral_Summaries_V6_EN.docx'


def add_table(doc, name):
    if name == 'relatedwork':
        table(doc, 'Table 1. Positioning relative to the closest lines.',
              ['Line', 'Info', 'Retained data', 'Result type', 'Shared budget',
               'Spectral compat.', 'Class-sharp', 'Boundary'],
              [['Robust RIS/beamforming [1],[4],[9]', 'full', 'matrices',
                'exact / conservative', 'yes (L1)', 'no', 'no', 'no'],
               ['Trust region / S-lemma [2],[3]', 'full', 'matrices',
                'exact', '—', '—', '—', 'hard case'],
               ['Probabilistic SINR bounds', 'distill.', 'moments',
                'conservative', 'no', 'no', 'no', 'no'],
               ['Buzano-type inequalities [5]', 'partial', 'norms',
                'two-sided inequality', 'no', 'partial', 'no', 'no'],
               ['PSD completion [6]', 'partial', 'specified entries',
                'feasibility', '—', '—', '—', '—'],
               ['Contraction completion [7]', 'partial', 'operator values',
                'set description', '—', 'partial', '—', 'Parrott'],
               ['Optimal recovery [8]', 'partial', 'finite data',
                'worst-case radius', '—', '—', 'yes (framework)', '—'],
               ['This work', 'partial', 'six scalars',
                'safe + sharp envelope', 'yes (L2)', 'yes', 'yes', 'yes']],
              [1.7, .55, 1.0, 1.05, .62, .72, .72, .55])
    elif name == 'boundary':
        table(doc, 'Table 2. Explicit boundary checks with a = 1 throughout.',
              ['p', 'a', 'I₀', 'J⊥', 'c∥', 'Q∂', 'Reference q*'],
              [['2', '1', '1', '0', '1', '0.590000', '0.590000'],
               ['0.2', '1', '1', '0', '−0.8', '−1.202000', '−1.202000'],
               ['2', '1', '1.3', '0.830662', '1', '−0.181020', '−0.181020'],
               ['0.2', '1', '1.3', '0.830662', '−0.8', '−1.963482', '−1.963482']],
              [.45, .4, .5, .8, .5, 1.2, 1.42])
    elif name == 'sensitivity':
        rows = [['0.5', '0.06', 118, 2, 117, 118, 118, 118, 0],
                ['0.5', '0.15', 108, 12, 91, 102, 103, 102, 1],
                ['0.5', '0.30', 63, 57, 6, 26, 28, 21, 2],
                ['1', '0.06', 90, 30, 78, 81, 83, 90, 2],
                ['1', '0.15', 56, 64, 21, 34, 36, 49, 2],
                ['1', '0.30', 7, 113, 0, 0, 0, 0, 0],
                ['2', '0.06', 24, 96, 13, 18, 18, 22, 0],
                ['2', '0.15', 6, 114, 0, 1, 4, 6, 3],
                ['2', '0.30', 0, 120, 0, 0, 0, 0, 0],
                ['Total', '', 472, 608, 326, 380, 390, 408, 10]]
        table(doc, 'Table 3. Directly accepted units on the complete sensitivity grid, 256 cells.',
              ['γ', 'ρ', 'Feas.', 'Infeas.', 'D', 'Old Q', 'Q_G', 'L_quad', 'Added'],
              rows, [.45, .5, .62, .68, .45, .68, .65, .82, .52])
    elif name == 'validation':
        rows = [['0.5', '0.06', 54, 53, 53, 53, 54],
                ['0.5', '0.15', 44, 28, 36, 37, 40],
                ['0.5', '0.30', 13, 0, 4, 4, 5],
                ['1', '0.06', 29, 22, 23, 23, 28],
                ['1', '0.15', 15, 5, 6, 6, 8],
                ['1', '0.30', 0, 0, 0, 0, 0],
                ['2', '0.06', 4, 3, 3, 3, 3],
                ['2', '0.15', 0, 0, 0, 0, 0],
                ['2', '0.30', 0, 0, 0, 0, 0],
                ['Total', '', 159, 111, 125, 126, 138]]
        table(doc, 'Table 4. Preregistered validation batch (seeds 211–215, 540 units, 256 cells).',
              ['γ', 'ρ', 'Feasible', 'D', 'Old Q', 'Q_G', 'L_quad'],
              rows, [.45, .5, .8, .5, .68, .65, .82])
    elif name == 'resolution':
        table(doc, 'Table 5. Interval resolution ablation (accepted units).',
              ['Cohort', 'Cells', 'Old Q', 'Q_G', 'Added', 'Feasible'],
              [['Exploratory', 256, 380, 390, 10, 472],
               ['Exploratory', 1024, 380, 390, 10, 472],
               ['Exploratory', 4096, 381, 390, 9, 472],
               ['Validation', 256, 125, 126, 1, 159],
               ['Validation', 1024, 125, 126, 1, 159],
               ['Validation', 4096, 125, 126, 1, 159]],
              [1.0, .6, .8, .8, .7, .8])
    elif name == 'fa':
        rows = [['0.5', '0.06', 59, 57, 59, 59, 59, 2],
                ['0.5', '0.15', 49, 40, 41, 44, 45, 1],
                ['0.5', '0.30', 23, 6, 7, 11, 13, 1],
                ['1', '0.06', 39, 29, 31, 38, 38, 2],
                ['1', '0.15', 22, 11, 11, 19, 19, 0],
                ['1', '0.30', 3, 1, 1, 0, 1, 0],
                ['2', '0.06', 11, 10, 10, 11, 11, 0],
                ['2', '0.15', 5, 2, 3, 4, 4, 1],
                ['2', '0.30', 0, 0, 0, 0, 0, 0],
                ['Total', '', 211, 156, 163, 186, 190, 7]]
        table(doc, 'Table 6. Realistic-link pool (seeds 301–305, distance-dependent path loss, '
                   'Rician/Rayleigh components; 540 units, 256 cells). ρ is the relative radius.',
              ['γ', 'ρ', 'Feasible', 'Old Q', 'Q_G', 'L_quad', 'Joint', 'Flips'],
              rows, [.45, .5, .75, .68, .65, .82, .7, .6])
    elif name == 'joint':
        table(doc, 'Table 7. Joint certification and selected pipeline (256 cells). '
                   'Confirmation entry: 202 is the common final accepted set; per-method '
                   'entries are reference-call totals.',
              ['Batch', 'Feasible', 'Q_G', 'L_quad', 'Joint', 'Speedup', 'Calls'],
              [['Exploratory (1080 units)', 472, 390, 408, 424, '2.12× (25/30)', '2946→731'],
               ['Validation (540 units)', 159, 126, 138, 139, '—', '—'],
               ['Realistic pool (540 units)', 211, 163, 186, 190, '—', '—'],
               ['Confirmation (540, fresh)', '202*', '471', '422', '394', '2.01× (11/15)', '1424→394']],
              [1.7, .7, .6, .7, .6, 1.15, .95])
    elif name == 'fb':
        table(doc, 'Table 8. Controlled scaling: dimension varies, all else fixed '
                   '(4 users, per-BS power 1, γ = 1, ρ = 0.15; medians over three seeds; '
                   'ms per user problem; pipeline seconds per seed).',
              ['Dim', 'Bound ms', 'Ref ms', 'Ref/Bound', 'Pure ref (s)',
               'L_quad + ref (s)', 'Lq→Q_G + ref (s)'],
              [['4', '0.041', '0.356', '8.7×', '0.0192', '0.0193', '0.0254'],
               ['8', '0.050', '0.368', '7.4×', '0.0378', '0.0165', '0.0236'],
               ['16', '0.044', '0.381', '8.7×', '0.0563', '0.0175', '0.0161'],
               ['32', '0.042', '0.435', '10.4×', '0.0859', '0.0310', '0.0310']],
              [.45, .8, .8, .85, .95, 1.1, 1.15])
    elif name == 'sdp':
        table(doc, 'Table 9. S-procedure SDP baseline on the 2160-problem validation batch.',
              ['Quantity', 'Value'],
              [['L ≤ q* (lower-bound safety), max excess', '−1.2e-13 (never above)'],
               ['U ≥ q* (witness validity), max deficit', '1.3e-12 (fp noise)'],
               ['Ordering L ≤ U, violations', '0'],
               ['False acceptance / false rejection', '0 / 0'],
               ['Witnesses inside the ball', '2160/2160 (956 projected)'],
               ['Decisions', '1247 feas., 913 infeas., 0 unresolved'],
               ['Witness quality U − q* (913 cases)', 'median 6.7e-6, max 1.6e-4'],
               ['Max |t − q*|', '5.6e-8'],
               ['Per solve: build / solver / wall', '0.33 / 1.6 / 5.7 ms']],
              [3.4, 2.3])
    elif name == 'timing':
        table(doc, 'Table 10. End-to-end seconds against the original 200-iteration reference '
                   '(supplementary); total reference calls across ten seeds in parentheses.',
              ['γ', 'Reference', 'D + fallback', 'Old Q + fallback', 'Q_G + fallback', 'L_quad + fallback'],
              [['0.5', '0.1787\n(1321)', '0.0375\n(244)', '0.0313\n(160)', '0.0316\n(147)', '0.0248\n(164)'],
               ['1', '0.1396\n(1005)', '0.0564\n(379)', '0.0516\n(306)', '0.0511\n(297)', '0.0429\n(285)'],
               ['2', '0.0874\n(620)', '0.0607\n(410)', '0.0612\n(384)', '0.0640\n(377)', '0.0536\n(353)']],
              [.35, 1., 1., 1., 1., .92])
    elif name == 'toltiming':
        table(doc, 'Table 11. Same final decision task against the validated tolerance-stopped '
                   'reference (median seconds; total calls in parentheses).',
              ['Pipeline', '200-iteration reference', 'Tolerance-stopped reference'],
              [['Pure reference', '0.1496  (1319)', '0.0411  (1319)'],
               ['New bound + fallback', '0.0750  (496)', '0.0323  (496)']],
              [1.8, 2.0, 2.0])
    else:
        raise ValueError(name)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.2677)
    sec.page_height = Inches(11.6929)
    sec.top_margin = sec.bottom_margin = Inches(.85)
    sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(5)
    for name, size in [('Title', 18), ('Heading 1', 13), ('Heading 2', 11.5), ('Caption', 9)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        if name != 'Caption':
            style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name.startswith('Heading') else 0)
        style.paragraph_format.space_after = Pt(5)
    equation = doc.styles.add_style('Equation', 1)
    equation.base_style = normal
    equation.font.name = 'Cambria Math'
    equation.font.size = Pt(10.5)
    equation.paragraph_format.space_before = Pt(3)
    equation.paragraph_format.space_after = Pt(6)
    equation.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_ALIGN_PARAGRAPH.RIGHT)
    settings = doc.settings.element
    mp = el('mathPr')
    mp.append(el('mathFont', val='Cambria Math'))
    settings.append(mp)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    footer._p.append(fld)

    for raw in (ROOT / 'paper_v6.md').read_text(encoding='utf-8').splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
            continue
        if line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
            continue
        if line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
            continue
        if line.startswith('EQ '):
            label, expr = line[3:].split(':', 1)
            p = doc.add_paragraph(style='Equation')
            p._p.append(math(expr.strip().rstrip('.')))
            p.add_run('\t' + label)
            continue
        if line.startswith('TABLE '):
            add_table(doc, line[6:])
            continue
        if line == 'FIGURE spectral':
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(PROJECT / 'results_review/figures/fig2_spectral_scan.png'),
                                    width=Inches(6.2))
            doc.add_paragraph('Figure 1. Constructed spectral scan. Approximate envelopes are '
                              'diagnostics; the dashed interval lower bound is used for '
                              'certification. The family changes the nominal channel norm.',
                              'Caption')
            continue
        p = doc.add_paragraph()
        add_inline(p, line)

    doc.core_properties.title = 'Sharp Robust QoS Certification from Six-Scalar Spectral Summaries'
    doc.core_properties.subject = ('English V6: information-class formulation, validation, '
                                   'complementary certificates, corrected SDP baseline')
    doc.core_properties.version = '6.0'
    for tree in (doc.styles.element, doc.element):
        for border in list(tree.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        for fonts in tree.iter(qn('w:rFonts')):
            for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
                fonts.attrib.pop(qn('w:' + attr), None)
    for name in ('Title', 'Heading 1', 'Heading 2', 'Caption'):
        doc.styles[name].font.name = 'Times New Roman'
    ppr = doc.paragraphs[0]._p.get_or_add_pPr()
    no_border = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        child = OxmlElement('w:' + side)
        child.set(qn('w:val'), 'nil')
        no_border.append(child)
    ppr.append(no_border)
    doc.save(OUTPUT)
    print(OUTPUT)
    import zipfile
    from lxml import etree
    with zipfile.ZipFile(OUTPUT) as z:
        xml = etree.fromstring(z.read('word/document.xml'))
        ns = {'m': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
              'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        print('native equations', len(xml.findall('.//m:oMath', ns)),
              'tables', len(xml.findall('.//w:tbl', ns)))


if __name__ == '__main__':
    main()
